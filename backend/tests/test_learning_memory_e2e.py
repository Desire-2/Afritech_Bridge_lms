"""End-to-end: learning memory is recorded only after an instructor review.

Grades a synthetic Excel submission through the real pipeline
(ExcelGradingService.grade_submission → analyzers → requirement evaluation →
persistence) against an isolated in-memory SQLite database, then asserts:

  1. auto-grading alone records ZERO GradingExperience rows,
  2. the auto-approval helper still records nothing,
  3. an instructor review (approve/override) records exactly one experience,
  4. learning insights then surface requirement-level memory (spec 18-20).

The file download layer is monkeypatched so the test is hermetic; every
other stage (parse, analyzers, evaluator, rubric cache, spec persistence,
learning) runs for real.
"""

import json
from io import BytesIO

import pytest
from flask import Flask, current_app

from src.models.user_models import db, User, Role
from src.models.course_models import Course, Module, Assignment, AssignmentSubmission
import src.models.course_application  # noqa: F401  (register referenced course_applications table)
import src.models.grading_models  # noqa: F401  (register rubrics tables for create_all)
from src.models.excel_grading_models import (
    ExcelGradingResult,
    GradingExperience,
    GeneratedRubric,
    AssignmentAssessmentSpec,
)


ASSIGNMENT_6_INSTRUCTIONS = (
    "1. Use HLOOKUP to find September revenue in E1.\n"
    "2. Use SUMIF for North region revenue in E2.\n"
    "3. Use COUNTIF to count transactions above 4000 in E3.\n"
    "4. Use AVERAGEIF for Software average in E4.\n"
    "5. Use wildcard COUNTIF for products ending in top in E5."
)


# ─── Fixtures ─────────────────────────────────────────────────

@pytest.fixture
def app():
    """Isolated Flask app bound to an in-memory SQLite database."""
    flask_app = Flask(__name__)
    flask_app.config.update(
        TESTING=True,
        SQLALCHEMY_DATABASE_URI="sqlite:///:memory:",
    )
    db.init_app(flask_app)
    return flask_app


@pytest.fixture
def db_session(app):
    """Fresh in-memory schema; the app context stays open for the whole test."""
    with app.app_context():
        db.create_all()
        yield db.session
        db.session.remove()
        db.drop_all()


@pytest.fixture
def seeded(app, db_session):
    """Course 1-style assignment + submission with users and roles."""
    instructor_role = Role(name="instructor")
    student_role = Role(name="student")
    db_session.add_all([instructor_role, student_role])
    db_session.flush()

    instructor = User(
        username="instructor_e2e", email="instructor_e2e@test.local",
        password_hash="x", role_id=instructor_role.id,
    )
    student = User(
        username="student_e2e", email="student_e2e@test.local",
        password_hash="x", role_id=student_role.id,
    )
    db_session.add_all([instructor, student])
    db_session.flush()

    course = Course(
        title="Excel Mastery: From Beginner to Data Expert",
        description="Advanced Microsoft Excel skills for data professionals.",
        instructor_id=instructor.id,
    )
    db_session.add(course)
    db_session.flush()

    module = Module(title="Aggregation Essentials", course_id=course.id, order=2)
    db_session.add(module)
    db_session.flush()

    assignment = Assignment(
        title="Applying Essential Arithmetic and Aggregation Skills",
        description="Use lookup and conditional aggregation functions.",
        instructions=ASSIGNMENT_6_INSTRUCTIONS,
        course_id=course.id,
        module_id=module.id,
        instructor_id=instructor.id,
        points_possible=100,
    )
    db_session.add(assignment)
    db_session.flush()

    submission = AssignmentSubmission(
        assignment_id=assignment.id,
        student_id=student.id,
        file_url=json.dumps([{
            "url": "https://example.invalid/submission.xlsx",
            "filename": "submission.xlsx",
            "original_filename": "submission.xlsx",
        }]),
    )
    db_session.add(submission)
    db_session.commit()

    return {
        "course": course,
        "module": module,
        "assignment": assignment,
        "submission": submission,
        "instructor": instructor,
        "student": student,
    }


def _synthetic_workbook() -> bytes:
    """A perfect Assignment-6 submission: one criterion per required cell."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Category"
    ws["B1"] = "Amount"
    ws["A2"] = "North"
    ws["B2"] = 100
    ws["E1"] = "=HLOOKUP(A1,A2:E5,2,FALSE)"
    ws["E2"] = '=SUMIF(A:A,"North",B:B)'
    ws["E3"] = '=COUNTIF(B:B,">4000")'
    ws["E4"] = '=AVERAGEIF(A:A,"Software",B:B)'
    ws["E5"] = '=COUNTIF(A:A,"*top")'
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _grade(seeded, monkeypatch):
    """Run the full grading pipeline and return (result, result_row)."""
    from src.services.excel_grading import ExcelGradingService

    submission = seeded["submission"]
    service = ExcelGradingService()
    monkeypatch.setattr(service, "_download_file", lambda file_info: _synthetic_workbook())

    result = service.grade_submission(submission.id, "assignment", force=True)
    assert result["status"] == "completed", result

    result_row = ExcelGradingResult.query.filter_by(
        assignment_submission_id=submission.id,
    ).first()
    assert result_row is not None
    return result, result_row


# ─── Tests ────────────────────────────────────────────────────

def test_learning_memory_recorded_only_after_instructor_approve(seeded, db_session, monkeypatch):
    from src.services.excel_grading.learning_engine import LearningEngine
    from src.services.excel_grading.auto_grader import _auto_approve_grading_result

    submission = seeded["submission"]
    assignment = seeded["assignment"]
    course = seeded["course"]
    module = seeded["module"]

    # 1) Full pipeline grade — requirement-level, no manual review needed.
    # grade_submission returns the persisted result's to_dict (exactly what the
    # sync route feeds to the auto-approval helper in production).
    result, result_row = _grade(seeded, monkeypatch)
    grade_data = result["result"]
    assert grade_data["analysis_data"]["requirement_results"]
    assert grade_data["requirements_count"] == 5
    assert grade_data["manual_review_required"] is False
    assert round(grade_data["total_score"] / grade_data["max_score"] * 100, 1) == 100

    # 2) Auto-grading must NOT write learning memory.
    assert GradingExperience.query.count() == 0
    assert GradingExperience.query.filter_by(assignment_id=assignment.id).count() == 0

    # 2b) The rubric-cache path ran and persisted the generated rubric.
    assert GeneratedRubric.query.count() == 1

    # 3) The versioned assessment contract IS persisted by the pipeline.
    spec = AssignmentAssessmentSpec.query.filter_by(assignment_id=assignment.id).first()
    assert spec is not None
    assert spec.engine_version
    assert spec.source_hash == grade_data["analysis_data"]["assessment_spec"]["source_hash"]
    assert spec.approved is False

    # The rubric/result metadata columns are populated by the pipeline.
    assert result_row.assessment_spec_version
    assert result_row.assessment_spec_hash == spec.source_hash
    assert result_row.requirements_satisfied == 5
    assert result_row.analyzers_used

    # 4) Even the auto-approval helper records no learning experience.
    eligible = _auto_approve_grading_result(
        current_app._get_current_object(), submission.id, "assignment", grade_data,
    )
    assert eligible is True
    assert GradingExperience.query.count() == 0
    # It does mark the result as reviewed/approved, though.
    assert result_row.instructor_reviewed is True

    # 5) Instructor review records exactly one experience (mirrors review route).
    engine = LearningEngine()
    engine.record_grading_outcome(
        grading_result_id=result_row.id,
        assignment_id=assignment.id,
        course_id=course.id,
        module_id=module.id,
        ai_score=result_row.total_score,
        ai_max_score=result_row.max_score,
        instructor_action="approve",
        instructor_notes="Looks good",
        rubric_used=result_row.rubric_data or result_row.rubric_breakdown,
        requirements_used=(result_row.analysis_data or {}).get("assessment_spec"),
        analysis_summary=result_row.analysis_data,
    )
    assert GradingExperience.query.count() == 1
    experience = GradingExperience.query.first()
    assert experience.instructor_action == "approve"
    assert experience.requirements_snapshot is not None  # spec wired into learning
    assert experience.analysis_summary is not None

    # 6) Learning closure: insights now expose requirement-level memory.
    insights = engine.get_insights(
        assignment_id=assignment.id, course_id=course.id, module_id=module.id,
    )
    assert insights["sample_size"] == 1
    assert insights["memory_scope"] == "assignment"
    assert insights["requirement_patterns"], "requirement patterns should exist after review"


def test_override_records_experience_with_delta_after_review(seeded, db_session, monkeypatch):
    from src.services.excel_grading.learning_engine import LearningEngine
    from src.services.excel_grading import ExcelGradingService

    assignment = seeded["assignment"]
    course = seeded["course"]
    module = seeded["module"]

    _, result_row = _grade(seeded, monkeypatch)

    # A second student submission graded WITHOUT any review must still leave
    # the learning table untouched.
    second_student = User(
        username="student_e2e_2", email="student_e2e_2@test.local",
        password_hash="x",
        role_id=Role.query.filter_by(name="student").first().id,
    )
    db_session.add(second_student)
    db_session.commit()
    second_submission = AssignmentSubmission(
        assignment_id=assignment.id,
        student_id=second_student.id,
        file_url=json.dumps([{
            "url": "https://example.invalid/submission2.xlsx",
            "filename": "submission2.xlsx",
            "original_filename": "submission2.xlsx",
        }]),
    )
    db_session.add(second_submission)
    db_session.commit()

    service = ExcelGradingService()
    monkeypatch.setattr(service, "_download_file", lambda file_info: _synthetic_workbook())
    second_result = service.grade_submission(second_submission.id, "assignment", force=True)
    assert second_result["status"] == "completed"
    assert GradingExperience.query.count() == 0  # still nothing recorded
    # Spec persistence is idempotent: still exactly one versioned contract.
    assert AssignmentAssessmentSpec.query.count() == 1
    assert GeneratedRubric.query.count() == 1

    # Instructor OVERRIDES the grade → experience recorded with a delta.
    engine = LearningEngine()
    engine.record_grading_outcome(
        grading_result_id=result_row.id,
        assignment_id=assignment.id,
        course_id=course.id,
        module_id=module.id,
        ai_score=result_row.total_score,
        ai_max_score=result_row.max_score,
        instructor_action="override",
        instructor_score=70.0,
        instructor_notes="HLOOKUP answer cell was wrong",
        rubric_used=result_row.rubric_data or result_row.rubric_breakdown,
        requirements_used=(result_row.analysis_data or {}).get("assessment_spec"),
        analysis_summary=result_row.analysis_data,
    )
    assert GradingExperience.query.count() == 1
    experience = GradingExperience.query.first()
    assert experience.instructor_action == "override"
    assert experience.instructor_score == 70.0
    assert experience.score_delta == round(70.0 - result_row.total_score, 2)

    # The override is now visible as calibration signal + requirement memory.
    insights = engine.get_insights(
        assignment_id=assignment.id, course_id=course.id, module_id=module.id,
    )
    assert insights["sample_size"] == 1
    assert insights["avg_override_delta"] == experience.score_delta
    assert insights["requirement_patterns"]
    assert insights["memory_scope"] == "assignment"
