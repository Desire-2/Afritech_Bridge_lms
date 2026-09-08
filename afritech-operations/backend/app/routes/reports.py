from datetime import date, datetime, timedelta
from decimal import Decimal
import csv
import io

from flask import Blueprint, request, jsonify, Response

from ..extensions import db, limiter
from ..models import (
    ServiceTransaction, Expense, PayrollPeriod, Service, Employee,
    Attendance, DailyClosing, User,
)
from ..auth.auth import require_permission
from .helpers import json_error

bp = Blueprint('reports', __name__, url_prefix='/api/reports')


def _ensure_admin_or_report_viewer(user):
    if not user.has_permission('reports.view') and not user.is_super_admin:
        raise PermissionError('You do not have permission to view reports')


def _date_range():
    start = request.args.get('start')
    end = request.args.get('end')
    if not start:
        end = end and date.fromisoformat(end) or date.today()
        start = (end - timedelta(days=29))
    else:
        start = date.fromisoformat(start)
        end = end and date.fromisoformat(end) or date.today()
    return start, end


def _filtered_txns_query(start, end):
    q = ServiceTransaction.query.filter(
        ServiceTransaction.transaction_date >= start,
        ServiceTransaction.transaction_date <= end,
    )
    employee_id = request.args.get('employee_id', type=int)
    service_id = request.args.get('service_id', type=int)
    department = request.args.get('department')
    status = request.args.get('status', 'completed')
    if employee_id:
        q = q.filter_by(employee_id=employee_id)
    if service_id:
        q = q.filter_by(service_id=service_id)
    if department:
        q = q.join(Employee, Employee.id == ServiceTransaction.employee_id).filter(Employee.department_id == int(department))
    if status and status != 'all':
        q = q.filter_by(status=status)
    return q


def _row_vals(t):
    return {
        'transaction_number': t.transaction_number, 'date': t.transaction_date.isoformat(),
        'status': t.status, 'service': t.service_name, 'client': t.client.full_name if t.client else '',
        'employee': t.snapshot_employee_name(),
        'customer_price': float(t.customer_price), 'official_cost': float(t.official_cost),
        'gross_profit': float(t.gross_profit), 'commission_rate': float(t.commission_rate_used),
        'commission': float(t.commission_amount), 'company_profit': float(t.company_profit),
    }


def _sum(rows, key):
    return sum(r[key] for r in rows)


@bp.get('/summary')
@require_permission('reports.view')
def summary():
    start, end = _date_range()
    txns = _filtered_txns_query(start, end).all()

    revenue = sum(float(t.customer_price) for t in txns)
    costs = sum(float(t.official_cost) for t in txns)
    gross = sum(float(t.gross_profit) for t in txns)
    commission = sum(float(t.commission_amount) for t in txns)
    company = sum(float(t.company_profit) for t in txns)

    expenses_q = Expense.query.filter(Expense.expense_date >= start, Expense.expense_date <= end)
    if request.args.get('category'):
        expenses_q = expenses_q.filter_by(category=request.args.get('category'))
    expenses = sum(float(e.amount) for e in expenses_q.all() if e.status in ('approved', 'paid'))

    payroll = 0.0
    periods = PayrollPeriod.query.filter(PayrollPeriod.period_end >= start, PayrollPeriod.period_start <= end).all()
    for p in periods:
        if p.status in ('approved', 'paid'):
            payroll += sum(float(i.net_salary or 0) for i in p.items)

    return jsonify({
        'period': {'start': start.isoformat(), 'end': end.isoformat()},
        'transaction_count': len(txns),
        'revenue': round(revenue, 2),
        'service_costs': round(costs, 2),
        'gross_profit': round(gross, 2),
        'commissions': round(commission, 2),
        'company_profit': round(company, 2),
        'expenses': round(expenses, 2),
        'payroll': round(payroll, 2),
        'net_profit': round(company - expenses, 2),
    })


@bp.get('/revenue')
@require_permission('reports.view')
def revenue_report():
    start, end = _date_range()
    txns = _filtered_txns_query(start, end).all()
    rows = [_row_vals(t) for t in txns]
    return jsonify({
        'period': {'start': start.isoformat(), 'end': end.isoformat()},
        'total_revenue': round(_sum(rows, 'customer_price'), 2),
        'total_gross_profit': round(_sum(rows, 'gross_profit'), 2),
        'total_company_profit': round(_sum(rows, 'company_profit'), 2),
        'total_commission': round(_sum(rows, 'commission'), 2),
        'count': len(rows),
        'rows': rows,
    })


@bp.get('/by-service')
@require_permission('reports.view')
def by_service():
    start, end = _date_range()
    txns = _filtered_txns_query(start, end).all()
    agg = {}
    for t in txns:
        svc = t.service_name
        a = agg.setdefault(svc, {'service': svc, 'count': 0, 'revenue': 0, 'cost': 0,
                                 'gross_profit': 0, 'commission': 0, 'company_profit': 0})
        a['count'] += 1
        a['revenue'] += float(t.customer_price)
        a['cost'] += float(t.official_cost)
        a['gross_profit'] += float(t.gross_profit)
        a['commission'] += float(t.commission_amount)
        a['company_profit'] += float(t.company_profit)
    items = sorted(agg.values(), key=lambda x: x['revenue'], reverse=True)
    return jsonify({'items': items, 'total': len(items)})


@bp.get('/by-employee')
@require_permission('reports.view')
def by_employee():
    start, end = _date_range()
    txns = _filtered_txns_query(start, end).all()
    agg = {}
    for t in txns:
        name = t.snapshot_employee_name()
        a = agg.setdefault(name, {'employee': name, 'count': 0, 'revenue': 0, 'gross_profit': 0,
                                  'commission': 0, 'company_profit': 0})
        a['count'] += 1
        a['revenue'] += float(t.customer_price)
        a['gross_profit'] += float(t.gross_profit)
        a['commission'] += float(t.commission_amount)
        a['company_profit'] += float(t.company_profit)
    items = sorted(agg.values(), key=lambda x: x['company_profit'], reverse=True)
    return jsonify({'items': items})


@bp.get('/expenses')
@require_permission('reports.view')
def expense_report():
    start, end = _date_range()
    q = Expense.query.filter(Expense.expense_date >= start, Expense.expense_date <= end)
    category = request.args.get('category')
    status = request.args.get('status')
    if category:
        q = q.filter_by(category=category)
    if status:
        q = q.filter_by(status=status)
    expenses = q.order_by(Expense.expense_date.desc()).all()
    by_cat = {}
    total = 0
    for e in expenses:
        total += float(e.amount)
        c = by_cat.setdefault(e.category, {'category': e.category, 'count': 0, 'total': 0.0})
        c['count'] += 1
        c['total'] += float(e.amount)
    return jsonify({
        'total': round(total, 2),
        'count': len(expenses),
        'by_category': sorted(by_cat.values(), key=lambda x: x['total'], reverse=True),
        'rows': [e.to_dict() for e in expenses],
    })


@bp.get('/attendance')
@require_permission('reports.view')
def attendance_report():
    start, end = _date_range()
    q = Attendance.query.filter(Attendance.attendance_date >= start, Attendance.attendance_date <= end)
    employee_id = request.args.get('employee_id', type=int)
    if employee_id:
        q = q.filter_by(employee_id=employee_id)
    records = q.all()
    by_status = {}
    for r in records:
        s = by_status.setdefault(r.status, {'status': r.status, 'count': 0})
        s['count'] += 1
    by_employee = {}
    for r in records:
        name = r.employee.full_name if r.employee else '?'
        e = by_employee.setdefault(r.employee_id, {'employee_id': r.employee_id, 'employee': name,
                                                   'present': 0, 'late': 0, 'absent': 0, 'leave': 0, 'total_hours': 0.0})
        e[r.status] = e.get(r.status, 0) + 1
        e['total_hours'] += float(r.total_hours or 0)
    return jsonify({
        'total_records': len(records),
        'by_status': by_status,
        'by_employee': list(by_employee.values()),
    })


@bp.get('/cash-reconciliation')
@require_permission('reports.view')
def cash_recon_report():
    start, end = _date_range()
    q = DailyClosing.query.filter(DailyClosing.closing_date >= start, DailyClosing.closing_date <= end)
    closings = q.all()
    total_expected = sum(float(c.expected_cash or 0) for c in closings)
    total_actual = sum(float(c.actual_cash or 0) for c in closings)
    shortage_count = sum(1 for c in closings if c.reconciliation_class == 'shortage')
    overage_count = sum(1 for c in closings if c.reconciliation_class == 'overage')
    exact_count = sum(1 for c in closings if c.reconciliation_class == 'exact')
    return jsonify({
        'closings': [c.to_dict() for c in closings],
        'total_expected': round(total_expected, 2),
        'total_actual': round(total_actual, 2),
        'total_difference': round(total_actual - total_expected, 2),
        'shortage_count': shortage_count,
        'overage_count': overage_count,
        'exact_count': exact_count,
        'count': len(closings),
    })


@bp.get('/payroll')
@require_permission('reports.view')
def payroll_report():
    start, end = _date_range()
    periods = PayrollPeriod.query.filter(PayrollPeriod.period_start <= end, PayrollPeriod.period_end >= start).all()
    items = []
    for p in periods:
        for i in p.items:
            d = i.to_dict()
            d['period'] = p.name
            d['period_status'] = p.status
            items.append(d)
    return jsonify({
        'periods': [{'name': p.name, 'status': p.status,
                     'total_net': sum(float(i.net_salary or 0) for i in p.items)} for p in periods],
        'items': items,
        'total_net': round(sum(float(i.get('net_salary') or 0) for i in items if i.get('period_status') in ('approved', 'paid')), 2),
    })


@bp.get('/instructor-performance')
@require_permission('reports.view')
def instructor_performance_report():
    from ..models import PerformanceScore, PerformanceMetric
    start = request.args.get('start')
    end = request.args.get('end')
    q = PerformanceScore.query
    if start:
        q = q.filter(PerformanceScore.period_start >= date.fromisoformat(start))
    if end:
        q = q.filter(PerformanceScore.period_end <= date.fromisoformat(end))
    scores = q.order_by(PerformanceScore.updated_at.desc()).all()
    return jsonify({'scores': [s.to_dict() for s in scores]})


# ---------- export ----------
FIELDS = ['transaction_number', 'date', 'status', 'service', 'client', 'employee',
          'customer_price', 'official_cost', 'gross_profit', 'commission_rate',
          'commission', 'company_profit']


@bp.get('/export/csv')
@require_permission('reports.export')
def export_csv():
    start, end = _date_range()
    report = request.args.get('report', 'revenue')
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    if report == 'revenue' or report == 'transactions':
        txns = _filtered_txns_query(start, end).all()
        writer.writerow(FIELDS)
        for t in txns:
            r = _row_vals(t)
            writer.writerow([r[f] for f in FIELDS])
    elif report == 'expenses':
        q = Expense.query.filter(Expense.expense_date >= start, Expense.expense_date <= end)
        expenses = q.all()
        writer.writerow(['id', 'date', 'category', 'description', 'amount', 'payment_method', 'status'])
        for e in expenses:
            writer.writerow([e.id, e.expense_date.isoformat(), e.category, e.description,
                             float(e.amount), e.payment_method.name if e.payment_method else '', e.status])
    else:
        return json_error('Unknown report type', 400)

    out = buffer.getvalue()
    response = Response(out, mimetype='text/csv')
    response.headers['Content-Disposition'] = f'attachment; filename={report}_{start}_{end}.csv'
    return response


@bp.get('/export/excel')
@require_permission('reports.export')
def export_excel():
    start, end = _date_range()
    report = request.args.get('report', 'revenue')
    try:
        import openpyxl
        from openpyxl.utils import get_column_letter
        from openpyxl.styles import Font, PatternFill
    except Exception:
        return json_error('Excel export is unavailable', 500)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='1F4778', end_color='1F4778', fill_type='solid')

    if report in ('revenue', 'transactions'):
        txns = _filtered_txns_query(start, end).all()
        ws.append(FIELDS)
        for t in txns:
            r = _row_vals(t)
            ws.append([r[f] for f in FIELDS])
    elif report == 'by-service':
        items = by_service().get_json()['items']
        ws.append(['service', 'count', 'revenue', 'cost', 'gross_profit', 'commission', 'company_profit'])
        for it in items:
            ws.append([it['service'], it['count'], it['revenue'], it['cost'], it['gross_profit'],
                       it['commission'], it['company_profit']])
    elif report == 'expenses':
        q = Expense.query.filter(Expense.expense_date >= start, Expense.expense_date <= end)
        ws.append(['id', 'date', 'category', 'description', 'amount', 'status'])
        for e in q.all():
            ws.append([e.id, e.expense_date.isoformat(), e.category, e.description, float(e.amount), e.status])
    else:
        return json_error('Unknown report type', 400)

    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
    for column_cells in ws.columns:
        length = max(len(str(c.value or '')) for c in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(length + 2, 40)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = Response(buffer.getvalue(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response.headers['Content-Disposition'] = f'attachment; filename={report}_{start}_{end}.xlsx'
    return response


@bp.get('/export/transactions-pdf')
@require_permission('reports.export')
def export_pdf():
    start, end = _date_range()
    txns = _filtered_txns_query(start, end).all()
    rows = [_row_vals(t) for t in txns]
    total_rev = round(_sum(rows, 'customer_price'), 2)
    total_comm = round(_sum(rows, 'commission'), 2)
    total_profit = round(_sum(rows, 'company_profit'), 2)

    html = f"""<html><head><meta charset="utf-8"><style>
    body {{ font-family: DejaVu Sans, sans-serif; font-size: 11px; }}
    h2 {{ color: #1F4778; }} table {{ width: 100%; border-collapse: collapse; }}
    th, td {{ border: 1px solid #ccc; padding: 4px; text-align: left; }}
    th {{ background: #1F4778; color: #fff; }}
    </style></head><body>
    <h2>Business Operations - {report_title()}</h2>
    <p>Period: {start} to {end}</p>
    <table><tr><th>#</th><th>Service</th><th>Client</th><th>Employee</th><th>Revenue</th><th>Cost</th>
    <th>Gross</th><th>Commission</th><th>Profit</th></tr>"""

    for i, r in enumerate(rows, 1):
        html += f"<tr><td>{i}</td><td>{r['service']}</td><td>{r['client']}</td><td>{r['employee']}</td>" \
                f"<td>{r['customer_price']:,.0f}</td><td>{r['official_cost']:,.0f}</td>" \
                f"<td>{r['gross_profit']:,.0f}</td><td>{r['commission']:,.0f}</td><td>{r['company_profit']:,.0f}</td></tr>"
    html += f"""</table>
    <p><b>Total revenue:</b> {total_rev:,.0f} | <b>Total commission:</b> {total_comm:,.0f} |
    <b>Company profit:</b> {total_profit:,.0f}</p></body></html>"""

    try:
        import pdfkit
        pdf = pdfkit.from_string(html, False)
    except Exception:
        try:
            import weasyprint
            pdf = weasyprint.HTML(string=html).write_pdf()
        except Exception:
            pdf = _export_pdf_reportlab(rows, start, end)

    response = Response(pdf, mimetype='application/pdf')
    response.headers['Content-Disposition'] = f'attachment; filename=transactions_{start}_{end}.pdf'
    return response


def _export_pdf_reportlab(rows, start, end):
    """Pure-python fallback when pdfkit/weasyprint are unavailable."""
    from io import BytesIO
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (Paragraph, SimpleDocTemplate, Spacer,
                                    Table, TableStyle)

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), rightMargin=24,
                            leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    story = [Paragraph(f'Business Operations - {report_title()}', styles['Title']),
             Paragraph(f'Period: {start} to {end}', styles['Normal']),
             Spacer(1, 10)]
    header = ['#', 'Service', 'Client', 'Employee', 'Revenue', 'Cost',
              'Gross', 'Commission', 'Profit']
    data = [header]
    for i, r in enumerate(rows, 1):
        data.append([str(i), r['service'], r['client'], r['employee'],
                     f"{r['customer_price']:,.0f}", f"{r['official_cost']:,.0f}",
                     f"{r['gross_profit']:,.0f}", f"{r['commission']:,.0f}",
                     f"{r['company_profit']:,.0f}"])
    total_rev = round(_sum(rows, 'customer_price'), 2)
    total_comm = round(_sum(rows, 'commission'), 2)
    total_profit = round(_sum(rows, 'company_profit'), 2)
    data.append(['', '', '', 'TOTAL',
                 f'{total_rev:,.0f}', '',
                 f'{round(_sum(rows, "gross_profit"), 2):,.0f}',
                 f'{total_comm:,.0f}', f'{total_profit:,.0f}'])
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1F4778')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTNAME', (0, -1), (-1, -1), 'Helvetica-Bold'),
        ('BACKGROUND', (0, -1), (-1, -1), colors.HexColor('#E9EEF5')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(tbl)
    doc.build(story)
    return buf.getvalue()


def report_title():
    return request.args.get('title', 'Transactions Report')